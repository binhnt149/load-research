import csv
import json
from dataclasses import asdict, replace

import pytest
from openpyxl import Workbook
from test_pipeline import request_files

from load_research.application.inspection import inspect_customer
from load_research.application.models import SourceInput
from load_research.application.pipeline import run_study
from load_research.application.result_export import export_selection
from load_research.infrastructure.store import IntervalStore
from load_research.result_view import ResultSelection
from load_research.workspace import Workspace


def test_offset_timestamp_is_imported_and_original_is_auditable(tmp_path):
    request = request_files(tmp_path)
    source = tmp_path / "mdms.csv"
    source.write_text(
        source.read_text().replace(
            "M1,2024-02-01T01:00:00,2\n", "M1,2024-02-01T01:02:02,2\n"
        )
    )
    workspace = Workspace.create(tmp_path / "project")
    result = run_study(workspace, request)
    assert result["import"]["timestamp_adjustments"] == 1
    detail = inspect_customer(
        workspace, result["run_id"], "0001", "2024-02-01", "2024-02-01"
    )
    assert detail["rows"][2]["original_kwh"] == 2
    database = workspace.safe_path(f"cache/{result['import']['key']}.duckdb")
    with IntervalStore(database, read_only=True) as store:
        corrections = list(store.iter_timestamp_adjustments())
    assert len(corrections) == 1
    assert corrections[0]["original_timestamp"] == "2024-02-01T01:02:02"
    assert corrections[0]["aligned_timestamp"] == "2024-02-01T01:00:00"
    assert corrections[0]["offset_seconds"] == 122

    second = run_study(workspace, request)
    assert second["import"]["key"] == result["import"]["key"]
    assert second["import"]["timestamp_adjustments"] == 1


def test_floor_collision_with_different_measurements_is_rejected(tmp_path):
    request = request_files(tmp_path)
    request = replace(
        request, config=replace(request.config, collision_policy="reject")
    )
    with (tmp_path / "mdms.csv").open("a") as stream:
        stream.write("M1,2024-02-01T01:25:00,50\n")
    workspace = Workspace.create(tmp_path / "project")
    with pytest.raises(ValueError, match="conflicting duplicate"):
        run_study(workspace, request)


def test_nearest_timestamp_wins_with_complete_conflict_audit(tmp_path):
    request = request_files(tmp_path)
    source = tmp_path / "mdms.csv"
    source.write_text(
        source.read_text().replace(
            "M1,2024-02-01T01:30:00,2\n", "M1,2024-02-01T01:35:00,7\n"
        )
    )
    with source.open("a") as stream:
        stream.write("M1,2024-02-01T01:25:00,50\n")
        stream.write("M1,2024-02-01T01:50:00,9\n")
    workspace = Workspace.create(tmp_path / "project")
    result = run_study(workspace, request)
    detail = inspect_customer(
        workspace, result["run_id"], "0001", "2024-02-01", "2024-02-01"
    )
    assert detail["rows"][2]["original_kwh"] == 2
    assert detail["rows"][3]["original_kwh"] == 7
    assert result["import"]["timestamp_conflict_groups"] == 2
    database = workspace.safe_path(f"cache/{result['import']['key']}.duckdb")
    with IntervalStore(database, read_only=True) as store:
        audit = list(store.iter_timestamp_conflicts())
    assert len(audit) == 4
    assert sum(row["selected"] for row in audit) == 2
    assert {row["offset_seconds"] for row in audit if not row["selected"]} == {
        1200,
        1500,
    }
    output = export_selection(
        workspace,
        result,
        tmp_path / "conflict-report",
        {
            "view_selection": asdict(ResultSelection("2024-02")),
            "include_xlsx": False,
            "include_timestamp_adjustments": True,
        },
        lambda: None,
    )
    with (output / "timestamp-conflicts.csv").open(
        encoding="utf-8-sig", newline=""
    ) as stream:
        conflicts = list(csv.DictReader(stream))
    assert len(conflicts) == 4
    assert sum(row["Được chọn"] == "Có" for row in conflicts) == 2


def test_nearest_tie_with_different_values_is_rejected(tmp_path):
    request = request_files(tmp_path)
    source = tmp_path / "mdms.csv"
    source.write_text(
        source.read_text().replace(
            "M1,2024-02-01T01:30:00,2\n", "M1,2024-02-01T01:35:00,7\n"
        )
    )
    with source.open("a") as stream:
        stream.write("M1,2024-02-01T01:35:00,9\n")
    with pytest.raises(ValueError, match="conflicting duplicate"):
        run_study(Workspace.create(tmp_path / "project"), request)


def test_conflict_audit_does_not_cross_join_workbook_sheets(tmp_path):
    request = request_files(tmp_path)
    workbook = Workbook()
    for sheet, meter, stamp, value in [
        (workbook.active, "M1", "2024-02-01T00:00:00", 1),
        (workbook.create_sheet("second"), "M2", "2024-02-01T00:10:00", 3),
        (workbook.create_sheet("third"), "M1", "2024-02-01T00:02:00", 2),
    ]:
        sheet.append(["meter_id", "timestamp", "value"])
        sheet.append([meter, stamp, value])
    path = tmp_path / "multiple-sheets.xlsx"
    workbook.save(path)
    request = replace(
        request,
        mdms_sources=tuple(
            SourceInput(
                str(path), replace(request.mdms_sources[0].options, sheet=sheet)
            )
            for sheet in workbook.sheetnames
        ),
    )
    workspace = Workspace.create(tmp_path / "project")
    result = run_study(workspace, request)
    detail = inspect_customer(
        workspace, result["run_id"], "0001", "2024-02-01", "2024-02-01"
    )
    assert detail["rows"][0]["original_kwh"] == 1
    database = workspace.safe_path(f"cache/{result['import']['key']}.duckdb")
    with IntervalStore(database, read_only=True) as store:
        audit = list(store.iter_timestamp_conflicts())
    assert len(audit) == 2
    assert next(row for row in audit if row["selected"])["offset_seconds"] == 0


def test_legacy_strict_alignment_rejects_off_grid_row_before_conflict(tmp_path):
    request = request_files(tmp_path)
    options = replace(request.mdms_sources[0].options, timestamp_alignment="strict")
    request = replace(
        request,
        config=replace(request.config, collision_policy="reject"),
        mdms_sources=(replace(request.mdms_sources[0], options=options),),
    )
    with (tmp_path / "mdms.csv").open("a") as stream:
        stream.write("M1,2024-02-01T01:25:00,50\n")
    result = run_study(Workspace.create(tmp_path / "project"), request)
    assert result["import"]["issues_count"] == 1
    assert result["import"]["timestamp_conflict_groups"] == 0


def test_secondary_interval_boundary_adjustment_is_audited(tmp_path):
    request = request_files(tmp_path)
    source = tmp_path / "mdms.csv"
    with source.open(newline="") as stream:
        rows = list(csv.DictReader(stream))
    rows[2]["interval_end"] = "2024-02-01T01:02:02"
    with source.open("w", newline="") as stream:
        writer = csv.DictWriter(
            stream, fieldnames=["meter_id", "timestamp", "value", "interval_end"]
        )
        writer.writeheader()
        writer.writerows(rows)
    workspace = Workspace.create(tmp_path / "project")
    result = run_study(workspace, request)
    assert result["import"]["timestamp_adjustments"] == 1
    database = workspace.safe_path(f"cache/{result['import']['key']}.duckdb")
    with IntervalStore(database, read_only=True) as store:
        correction = next(store.iter_timestamp_adjustments())
    assert correction["original_interval_end"] == "2024-02-01T01:02:02"
    assert correction["original_timestamp"] == "2024-02-01T00:30:00"


def test_old_source_snapshot_keeps_strict_rejection(tmp_path):
    request = request_files(tmp_path)
    options = replace(request.mdms_sources[0].options, timestamp_alignment="strict")
    request = replace(
        request, mdms_sources=(replace(request.mdms_sources[0], options=options),)
    )
    source = tmp_path / "mdms.csv"
    source.write_text(
        source.read_text().replace(
            "M1,2024-02-01T01:00:00,2\n", "M1,2024-02-01T01:02:02,2\n"
        )
    )
    result = run_study(Workspace.create(tmp_path / "project"), request)
    assert result["import"]["issues_count"] == 1
    assert result["import"]["timestamp_adjustments"] == 0


def test_optional_export_contains_complete_timestamp_audit_and_manifest(tmp_path):
    request = request_files(tmp_path)
    source = tmp_path / "mdms.csv"
    source.write_text(
        source.read_text().replace(
            "M1,2024-02-01T01:00:00,2\n", "M1,2024-02-01T01:02:02,2\n"
        )
    )
    workspace = Workspace.create(tmp_path / "project")
    result = run_study(workspace, request)
    output = export_selection(
        workspace,
        result,
        tmp_path / "report",
        {
            "view_selection": asdict(ResultSelection("2024-02")),
            "include_xlsx": False,
            "include_timestamp_adjustments": True,
        },
        lambda: None,
    )
    with (output / "timestamp-adjustments.csv").open(
        encoding="utf-8-sig", newline=""
    ) as stream:
        rows = list(csv.DictReader(stream))
    assert len(rows) == 1
    assert rows[0]["Tệp nguồn"] == "mdms.csv"
    assert rows[0]["SHA-256 nguồn"] == result["sources"][1]["sha256"]
    assert rows[0]["Thời điểm nguồn"] == "2024-02-01T01:02:02"
    assert rows[0]["Thời điểm sau hiệu chỉnh"] == "2024-02-01T01:00:00"
    assert rows[0]["Độ lệch về sau so với mốc (giây)"] == "122.0"
    manifest = json.loads((output / "export-manifest.json").read_text())
    assert "timestamp-adjustments.csv" in {item["name"] for item in manifest["files"]}


def test_timestamp_audit_export_escapes_spreadsheet_formula_and_checks_cache(tmp_path):
    request = request_files(tmp_path)
    cmis, source = tmp_path / "cmis.csv", tmp_path / "mdms.csv"
    cmis.write_text(cmis.read_text().replace("M1", "=M1"))
    source.write_text(
        source.read_text()
        .replace("M1,", "=M1,")
        .replace("2024-02-01T01:00:00,2\n", "2024-02-01T01:02:02,2\n")
    )
    workspace = Workspace.create(tmp_path / "project")
    result = run_study(workspace, request)
    options = {
        "view_selection": asdict(ResultSelection("2024-02")),
        "include_xlsx": False,
        "include_timestamp_adjustments": True,
    }
    output = export_selection(
        workspace, result, tmp_path / "report", options, lambda: None
    )
    with (output / "timestamp-adjustments.csv").open(
        encoding="utf-8-sig", newline=""
    ) as stream:
        row = next(csv.DictReader(stream))
    assert row["Mã công tơ"] == "'=M1"
    result["import"]["database_hash"] = "0" * 64
    with pytest.raises(ValueError, match="cache is missing or changed"):
        export_selection(
            workspace, result, tmp_path / "rejected", options, lambda: None
        )
    assert not (tmp_path / "rejected").exists()
