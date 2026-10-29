import csv
from dataclasses import replace

import numpy as np
import openpyxl
import pytest
from test_pipeline import request_files, write_csv

from load_research.application.daily import load_daily
from load_research.application.pipeline import run_study
from load_research.application.result_export import export_selection
from load_research.daily import expand_daily
from load_research.result_view import (
    ResultSelection,
    build_daily_view,
    build_result_view,
)
from load_research.workspace import Workspace


def setup_stratified_request(tmp_path, *, missing_group_sample=False):
    request = request_files(tmp_path)
    write_csv(
        tmp_path / "cmis.csv",
        ["customer_id", "industry_code", "meter_id", "annual_energy_kwh"],
        [
            ["0001", "01", "M1", 5000],
            ["0002", "01", "M2", 6000 if missing_group_sample else 20000],
        ],
    )
    path = tmp_path / "strata.csv"
    write_csv(
        path,
        [
            "Mã ngành nghề",
            "Nhóm sản lượng",
            "Số lượng khách hàng",
            "Tổng Sản lượng",
            "Trung bình sản lượng\n(Tổng sản lượng/số lượng khách hàng)",
            "Độ lệch chuẩn",
            "Độ lệch chuẩn/Trung bình sản lượng",
        ],
        [
            ["01", "A", 10, 50000, 5000, 1000, 0.2],
            ["01", "B", 20, 150000, 7500, 2000, 0.267],
        ],
    )
    return replace(request, strata_paths=(str(path),))


def test_industry_expansion_sums_separate_strata_and_preserves_monthly_energy(tmp_path):
    request = setup_stratified_request(tmp_path)
    workspace = Workspace.create(tmp_path / "project")
    result = run_study(workspace, request)
    cell = result["cells"][0]
    assert cell["stratification"]["method"] == "reference_energy_share"
    assert cell["result"]["status"] == "valid_with_warnings"
    assert (
        "reference_share_assumes_monthly_stability"
        in cell["result"]["diagnostics"]["warnings"]
    )
    strata = cell["stratification"]["strata"]
    assert strata["A"]["allocated_billing_kwh"] == pytest.approx(25000)
    assert strata["B"]["allocated_billing_kwh"] == pytest.approx(75000)
    assert strata["A"]["cohort_size"] == strata["B"]["cohort_size"] == 1
    assert {member["consumption_stratum"] for member in cell["members"]} == {"A", "B"}
    for kind, curve in cell["result"]["energy"].items():
        assert np.allclose(
            curve,
            np.array(strata["A"]["result"]["energy"][kind])
            + np.array(strata["B"]["result"]["energy"][kind]),
        )
    assert sum(
        cell["day_counts"][kind] * sum(curve)
        for kind, curve in cell["result"]["energy"].items()
    ) == pytest.approx(100000)
    assert result["request"]["strata_paths"] == [str(tmp_path / "strata.csv")]
    snapshot = workspace.load_artifact(result["run_id"], "stratum-config.json")
    assert snapshot["rows"][0]["population_count"] == 10
    assert len(snapshot["source_sha256"]) == 64
    view = build_result_view(result, ResultSelection("2024-02"))
    strata_series = [row for row in view["series"] if row["kind"] == "stratum"]
    assert {row["stratum"] for row in strata_series} == {"A", "B"}
    assert all(row["values"] is not None for row in strata_series)
    daily = load_daily(workspace, result, "2024-02", ("01",))["01"]
    assert set(daily["strata"]) == {"A", "B"}
    expanded = expand_daily(
        cell, daily, ResultSelection("2024-02", unit="kWh"), method="variation"
    )
    assert expanded["monthly_energy_kwh"] == pytest.approx(100000)
    daily_view = build_daily_view(
        result,
        ResultSelection("2024-02", unit="kWh"),
        {"01": daily},
        "2024-02-01",
        method="variation",
    )
    assert {row["stratum"] for row in daily_view["strata_factors"]} == {"A", "B"}
    assert all(row["daily_factor"] > 0 for row in daily_view["strata_factors"])
    output = export_selection(
        workspace,
        result,
        tmp_path / "export",
        {"view_selection": {"period": "2024-02"}, "include_xlsx": True},
        lambda: None,
    )
    assert (output / "strata-factors.csv").is_file()
    book = openpyxl.load_workbook(
        output / "profiles.xlsx", read_only=True, data_only=True
    )
    assert "Suy rộng nhóm sản lượng" in book.sheetnames
    assert book["Suy rộng nhóm sản lượng"].max_row == 3
    book.close()
    with (output / "profiles.csv").open(encoding="utf-8-sig", newline="") as stream:
        exported = list(csv.DictReader(stream))
    assert {row["Nhóm sản lượng"] for row in exported if row["Nhóm sản lượng"]} == {
        "A",
        "B",
    }


def test_positive_stratum_without_samples_blocks_complete_industry(tmp_path):
    request = setup_stratified_request(tmp_path, missing_group_sample=True)
    result = run_study(Workspace.create(tmp_path / "project"), request)
    cell = result["cells"][0]
    assert (
        cell["stratification"]["strata"]["B"]["result"]["status"]
        == "no_eligible_samples"
    )
    assert cell["result"]["energy"] == {}
    assert result["status"] == "partial"
    view = build_result_view(result, ResultSelection("2024-02"))
    assert next(row for row in view["series"] if row.get("stratum") == "A")["values"]
    assert (
        next(row for row in view["series"] if row.get("stratum") == "B")["values"]
        is None
    )


def test_sample_group_not_in_population_configuration_blocks_industry(tmp_path):
    request = setup_stratified_request(tmp_path)
    cmis = tmp_path / "cmis.csv"
    cmis.write_text(cmis.read_text().replace(",20000", ",200000"))
    path = tmp_path / "strata.csv"
    text = path.read_text()
    path.write_text(text[: text.index("01,B,")])
    result = run_study(Workspace.create(tmp_path / "project"), request)
    cell = result["cells"][0]
    assert cell["result"]["status"] == "missing_stratum_config"
    assert cell["result"]["energy"] == {}


def test_customer_below_501_kwh_is_reported_as_outside_scope(tmp_path):
    request = setup_stratified_request(tmp_path)
    cmis = tmp_path / "cmis.csv"
    cmis.write_text(cmis.read_text().replace(",5000", ",500"))
    result = run_study(Workspace.create(tmp_path / "project"), request)
    assert result["quality"][0]["excluded"]["0001"] == "annual_below_minimum"


def test_population_industry_without_samples_or_billing_blocks_organization(tmp_path):
    request = setup_stratified_request(tmp_path)
    with (tmp_path / "strata.csv").open("a") as stream:
        stream.write("02,A,10,50000,5000,1000,0.2\n")
    result = run_study(Workspace.create(tmp_path / "project"), request)
    assert set(result["industry_universe"]["2024-02"]) == {"01", "02"}
    view = build_result_view(result, ResultSelection("2024-02", level="organization"))
    assert not view["complete"]
    assert any(row["industry_code"] == "02" for row in view["missing"])
