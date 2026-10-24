import openpyxl
import pyarrow as pa
import pyarrow.parquet as pq
import pytest


def test_preview_stops_after_a_bounded_lookahead_and_keeps_source_headers(tmp_path):
    from load_research.importers.tables import read_preview

    path = tmp_path / "data.csv"
    path.write_text("customer_id,_custom\n0001,a\n0002,b\n0003,c\nbad,row,extra\n")
    preview = read_preview(path, limit=2)
    assert preview.columns == ("customer_id", "_custom")
    assert preview.rows == [
        {"customer_id": "0001", "_custom": "a"},
        {"customer_id": "0002", "_custom": "b"},
    ]
    assert preview.has_more


def test_preview_of_a_header_only_template_keeps_the_columns(tmp_path):
    from load_research.importers.tables import read_preview

    path = tmp_path / "empty.csv"
    path.write_text("customer_id,industry_code\n")
    preview = read_preview(path)
    assert preview.columns == ("customer_id", "industry_code")
    assert preview.rows == []
    assert not preview.has_more


def test_preview_supports_excel_sheet_csv_delimiter_and_parquet(tmp_path):
    from load_research.importers.tables import excel_sheets, read_preview

    path = tmp_path / "source.csv"
    path.write_text("customer_id;industry_code\n0001;01\n")
    assert read_preview(path, delimiter=";").rows[0]["customer_id"] == "0001"
    book = openpyxl.Workbook()
    book.active.title = "Giới thiệu"
    book.active.append(["Thông tin"])
    data = book.create_sheet("Dữ liệu")
    data.append(["customer_id", "industry_code"])
    data.append(["0001", "01"])
    xlsx = tmp_path / "source.xlsx"
    book.save(xlsx)
    assert excel_sheets(xlsx) == ("Giới thiệu", "Dữ liệu")
    assert read_preview(xlsx, sheet="Dữ liệu").rows[0]["customer_id"] == "0001"
    parquet = tmp_path / "source.parquet"
    pq.write_table(pa.table({"customer_id": ["0001"]}), parquet)
    assert read_preview(parquet).rows == [{"customer_id": "0001"}]


@pytest.mark.parametrize("limit", [0, -1, 1001, True])
def test_preview_rejects_invalid_limits(tmp_path, limit):
    from load_research.importers.tables import read_preview

    path = tmp_path / "source.csv"
    path.write_text("id\n0001\n")
    with pytest.raises(ValueError, match="limit"):
        read_preview(path, limit=limit)
