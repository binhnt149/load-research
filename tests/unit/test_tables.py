import csv

import openpyxl
import pyarrow as pa
import pyarrow.parquet as pq
import pytest

from load_research.importers.tables import read_batches, read_rows


def test_csv_mapping_preserves_customer_id_and_reads_in_batches(tmp_path):
    path = tmp_path / "khách hàng.csv"
    with path.open("w", newline="", encoding="utf-8-sig") as stream:
        writer = csv.writer(stream)
        writer.writerow(["Mã KH", "Ngành"])
        writer.writerows([["000123", "01"], ["000124", "02"]])
    batches = list(
        read_batches(
            path,
            mapping={"customer_id": "Mã KH", "industry_code": "Ngành"},
            batch_size=1,
        )
    )
    assert len(batches) == 2
    assert batches[0][0]["customer_id"] == "000123"
    assert batches[1][0]["industry_code"] == "02"


def test_xlsx_and_parquet_read_real_files(tmp_path):
    workbook = openpyxl.Workbook()
    workbook.active.append(["customer_id", "industry_code"])
    workbook.active.append(["000123", "01"])
    xlsx = tmp_path / "cmis.xlsx"
    workbook.save(xlsx)
    parquet = tmp_path / "cmis.parquet"
    pq.write_table(
        pa.table({"customer_id": ["000123"], "industry_code": ["01"]}), parquet
    )
    assert list(read_rows(xlsx))[0]["customer_id"] == "000123"
    assert list(read_rows(parquet))[0]["customer_id"] == "000123"


def test_uncached_excel_formula_is_not_executed_or_treated_as_a_value(tmp_path):
    workbook = openpyxl.Workbook()
    workbook.active.append(["customer_id"])
    workbook.active.append(['=CONCAT("000",123)'])
    path = tmp_path / "formula.xlsx"
    workbook.save(path)
    with pytest.raises(ValueError, match="formula"):
        list(read_rows(path))


def test_ambiguous_headers_and_missing_mapping_are_rejected(tmp_path):
    path = tmp_path / "bad.csv"
    path.write_text("id,id\n1,2\n")
    with pytest.raises(ValueError, match="header"):
        list(read_rows(path))
    path.write_text("id\n1\n")
    with pytest.raises(ValueError, match="mapping"):
        list(read_rows(path, mapping={"customer_id": "missing"}))
