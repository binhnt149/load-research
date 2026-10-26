import csv
import json

import numpy as np
import pytest
from matplotlib.backends.backend_agg import FigureCanvasAgg
from test_result_aggregation import study

from load_research.result_reporting import export_views, render_view
from load_research.result_view import (
    ResultSelection,
    build_daily_view,
    build_result_view,
)


def test_aggregate_chart_and_report_share_the_same_values(tmp_path):
    result = study()
    view = build_result_view(result, ResultSelection("2024-04", ("01", "02")))
    figure = render_view(view)
    FigureCanvasAgg(figure).draw()
    assert figure.axes[0].lines[0].get_ydata() == pytest.approx([32] * 48)
    destination = export_views(
        result, [view], tmp_path / "report", chart_formats=("png", "svg")
    )
    with (destination / "profiles.csv").open(encoding="utf-8-sig") as stream:
        rows = list(csv.DictReader(stream))
    overall = [row for row in rows if row["Loại đường biểu diễn"] == "Tổng hợp"]
    assert len(overall) == 48
    assert float(overall[0]["Giá trị"]) == 32
    manifest = json.loads((destination / "export-manifest.json").read_text())
    assert manifest["views"][0]["industry_codes"] == ["01", "02"]
    assert (destination / "profiles.xlsx").is_file()
    assert len(list((destination / "charts").iterdir())) == 2


def daily_data():
    return {
        code: {
            "dates": ["2024-04-01", "2024-04-06"],
            "day_types": ["weekday", "saturday"],
            "source": [[1] * 48, [2] * 48],
            "processed": [[2] * 48, [4] * 48],
            "eligible": [True, True],
            "cohort_size": 1,
        }
        for code in ("01", "02")
    }


def test_daily_view_adds_industries_and_keeps_both_sample_versions():
    view = build_daily_view(
        study(),
        ResultSelection("2024-04", ("01", "02"), interval_minutes=60, unit="kWh"),
        daily_data(),
        "2024-04-01",
        method="representative",
    )
    curves = {row["kind"]: row["values"] for row in view["series"]}
    assert curves["overall"] == [24] * 24
    assert curves["source"] == [4] * 24 and curves["processed"] == [8] * 24
    assert view["date"] == "2024-04-01" and view["method"] == "representative"
    assert view["energy_selected_kwh"] == 576


def test_daily_export_identifies_method_date_and_never_turns_gaps_into_zero(tmp_path):
    daily = daily_data()
    daily["01"]["source"][0][0] = None
    view = build_daily_view(
        study(),
        ResultSelection("2024-04", ("01", "02")),
        daily,
        "2024-04-01",
        method="variation",
    )
    assert view["series"][-2]["values"][0] is None
    assert np.isfinite(view["series"][0]["values"]).all()
    destination = export_views(study(), [view], tmp_path / "daily", include_xlsx=False)
    with (destination / "profiles.csv").open(encoding="utf-8-sig") as stream:
        rows = list(csv.DictReader(stream))
    source = [
        row for row in rows if row["Loại đường biểu diễn"] == "Mẫu nguồn trước VEE"
    ]
    assert source[0]["Giá trị"] == ""
    assert source[0]["Ngày"] == "2024-04-01"
    assert "biến thiên" in source[0]["Phương pháp theo ngày"]


def test_wide_report_has_24_slots_and_includes_quality_even_without_excel(tmp_path):
    result = study()
    result["quality"] = [{"period": "2024-04", "expected_intervals": 1440}]
    view = build_result_view(result, ResultSelection("2024-04", interval_minutes=60))
    output = export_views(
        result, [view], tmp_path / "wide", layout="wide", include_xlsx=False
    )
    with (output / "profiles.csv").open(encoding="utf-8-sig") as stream:
        rows = list(csv.DictReader(stream))
    assert len([key for key in rows[0] if key.startswith("Chu kỳ ")]) == 24
    assert rows[0]["Chu kỳ 01"] != ""
    assert "Số chu kỳ dự kiến" in (output / "quality-summary.csv").read_text(
        encoding="utf-8-sig"
    )
    assert (output / "balance-checks.csv").exists()


def test_report_strings_cannot_become_excel_formulas_and_cancellation_is_atomic(
    tmp_path,
):
    import openpyxl

    result = study()
    view = build_result_view(result, ResultSelection("2024-04", name="=1+1"))
    output = export_views(result, [view], tmp_path / "safe")
    with (output / "profiles.csv").open(encoding="utf-8-sig") as stream:
        rows = list(csv.DictReader(stream))
    assert (
        next(row for row in rows if row["Loại đường biểu diễn"] == "Tổng hợp")[
            "Tên đường"
        ]
        == "'=1+1"
    )
    book = openpyxl.load_workbook(output / "profiles.xlsx", read_only=True)
    try:
        assert not [
            cell
            for sheet in book
            for row in sheet
            for cell in row
            if cell.data_type == "f"
        ]
    finally:
        book.close()
    calls = 0

    def cancel():
        nonlocal calls
        calls += 1
        if calls > 1:
            raise RuntimeError("cancelled")

    with pytest.raises(RuntimeError, match="cancelled"):
        export_views(result, [view], tmp_path / "cancelled", check=cancel)
    assert not (tmp_path / "cancelled").exists()
    assert not list(tmp_path.glob(".load-research-report-*"))


def test_daily_unknown_organization_never_claims_a_complete_total():
    result = study()
    result.pop("industry_universe")
    view = build_daily_view(
        result,
        ResultSelection("2024-04", level="organization"),
        daily_data(),
        "2024-04-01",
        method="representative",
    )
    assert not view["complete"]
    assert view["series"][0]["values"] is None
    assert any(row["reason"] == "unknown_industry_universe" for row in view["missing"])


@pytest.mark.parametrize("method", ["representative", "variation"])
def test_daily_sample_scope_uses_actual_processed_sample_instead_of_expansion(method):
    result = study()
    view = build_daily_view(
        result,
        ResultSelection("2024-04", scope="sample_portfolio", unit="kWh"),
        daily_data(),
        "2024-04-01",
        method=method,
    )
    assert view["series"][0]["values"] == [4] * 48
    assert view["energy_selected_kwh"] == 192


def test_daily_date_must_match_selected_day_types():
    with pytest.raises(ValueError, match="day type"):
        build_daily_view(
            study(),
            ResultSelection("2024-04", day_types=("saturday",)),
            daily_data(),
            "2024-04-01",
            method="representative",
        )


def test_export_members_obey_selected_day_type_and_behavior_group(tmp_path):
    result = study()
    result["cells"][0]["members"] = [
        {"customer_id": "001", "day_type": "weekday", "cluster_id": "group_1"},
        {"customer_id": "002", "day_type": "weekday", "cluster_id": "group_2"},
        {"customer_id": "003", "day_type": "saturday", "cluster_id": "group_1"},
    ]
    view = build_result_view(result, ResultSelection("2024-04", ("01",), ("weekday",)))
    view["cluster_ids"] = ("group_1",)
    output = export_views(result, [view], tmp_path / "members", include_members=True)
    with (output / "cluster-members.csv").open(encoding="utf-8-sig") as stream:
        rows = list(csv.DictReader(stream))
    assert [row["Mã khách hàng"] for row in rows] == ["001"]
    assert rows[0]["Loại ngày"] == "Ngày làm việc"


def test_report_rows_and_factor_sheet_identify_period_scope_and_normalization(tmp_path):
    import openpyxl

    view = build_result_view(
        study(),
        ResultSelection("2024-04", scope="sample_portfolio", normalization="mean"),
    )
    output = export_views(study(), [view], tmp_path / "metadata")
    with (output / "profiles.csv").open(encoding="utf-8-sig") as stream:
        row = next(csv.DictReader(stream))
    assert row["Phạm vi phụ tải"] == "Phụ tải tập mẫu"
    assert row["Chuẩn hóa"] == "Trung bình bằng 1"
    assert "Ngày làm việc" in row["Các loại ngày đã chọn"]
    workbook = openpyxl.load_workbook(output / "profiles.xlsx", read_only=True)
    try:
        sheet = workbook["Hệ số và cân bằng"]
        headings = [cell.value for cell in next(sheet.rows)]
        assert "Năm-tháng" in headings
        assert sheet.cell(2, headings.index("Năm-tháng") + 1).value == "2024-04"
    finally:
        workbook.close()
