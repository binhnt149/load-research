import openpyxl
import pytest

from load_research.domain.config import SourceOptions
from load_research.importers.catalog import load_billing, load_catalog
from load_research.importers.metering import canonicalize
from load_research.importers.tables import read_rows

KEYS = (
    "cmis",
    "mdms_long",
    "mdms_wide_24",
    "mdms_wide_48",
    "mdms_wide_96",
    "billing",
    "calendar",
)


@pytest.mark.parametrize("key", KEYS)
@pytest.mark.parametrize("extension", ["csv", "xlsx"])
def test_every_template_can_be_read_by_the_real_importer(tmp_path, key, extension):
    from load_research.input_templates import write_template

    path = write_template(key, tmp_path / f"{key}.{extension}")
    rows = list(read_rows(path))
    assert rows
    if key == "cmis":
        assert rows[0]["customer_id"] == "000001"
        assert rows[0]["industry_code"] == "01"
        assert len(load_catalog(rows, "NPC").links) == 2
    elif key == "billing":
        assert load_billing(rows, "NPC")[("01", "2026-01")].energy_kwh == 100000
    elif key == "calendar":
        assert {row["day_type"] for row in rows} <= {
            "weekday",
            "saturday",
            "sunday",
            "holiday",
        }
    else:
        assert rows[0]["meter_id"] == "M000001"
        assert rows[0]["unit"] == "kWh"
    if extension == "xlsx":
        book = openpyxl.load_workbook(path, read_only=True, data_only=False)
        try:
            assert book.sheetnames == ["Dữ liệu", "Hướng dẫn"]
            assert all(
                cell.data_type != "f" for sheet in book for row in sheet for cell in row
            )
        finally:
            book.close()


@pytest.mark.parametrize("slots", [24, 48, 96])
def test_wide_templates_use_the_declared_resolution_without_changing_daily_energy(
    tmp_path, slots
):
    from load_research.input_templates import write_template

    cmis = write_template("cmis", tmp_path / "cmis.csv")
    mdms = write_template(f"mdms_wide_{slots}", tmp_path / "mdms.csv")
    catalog = load_catalog(read_rows(cmis), "NPC")
    issues = []
    records = list(
        canonicalize(
            read_rows(mdms),
            catalog,
            SourceOptions(
                values_already_scaled=True,
                layout="wide",
                interval_minutes=1440 // slots,
            ),
            issues=issues,
        )
    )
    assert issues == []
    assert len(records) == 2 * slots
    assert sum(record.energy_kwh for record in records) == 144


def test_template_export_never_overwrites_an_existing_source_and_cleans_staging(
    tmp_path,
):
    from load_research.input_templates import write_template

    path = tmp_path / "existing.xlsx"
    path.write_bytes(b"original customer data")
    with pytest.raises(FileExistsError):
        write_template("cmis", path)
    assert path.read_bytes() == b"original customer data"
    assert list(tmp_path.iterdir()) == [path]
    with pytest.raises(ValueError):
        write_template("../unknown", tmp_path / "new.csv")
    with pytest.raises(ValueError):
        write_template("cmis", tmp_path / "new.exe")
