"""Built-in, synthetic input examples and safe CSV/XLSX template export."""

import csv
import os
import tempfile
from dataclasses import dataclass
from pathlib import Path

import openpyxl
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

NOTICE = (
    "Dữ liệu minh hoạ. Thay toàn bộ dòng ví dụ bằng dữ liệu thực trước khi nghiên cứu."
)
FIELD_HELP = {
    "customer_id": "Mã khách hàng dạng chuỗi; giữ nguyên số 0 đầu.",
    "industry_code": "Mã ngành nghề dạng chuỗi, dùng thống nhất giữa CMIS và thương phẩm.",
    "industry_name": "Tên ngành nghề; nhất quán theo mã ngành.",
    "meter_id": "Mã công tơ ghép với CMIS; không ghép bằng tên khách hàng.",
    "organization_code": "Mã đơn vị, khớp phạm vi nghiên cứu.",
    "sampling_weight": "Trọng số dương đã được chốt; dùng 1 nếu không có trọng số riêng.",
    "timestamp": "ISO 8601 theo giờ địa phương; ví dụ dùng dấu đầu chu kỳ.",
    "value": "Điện năng kWh trong chu kỳ; thiếu để trống, số 0 là số đo hợp lệ.",
    "source_quality_flag": "actual = đo gốc; estimated = ước lượng; missing/invalid = thiếu.",
    "unit": "Ví dụ dùng kWh; phải khớp đơn vị khai báo của nguồn.",
    "direction": "import = điện năng tác dụng chiều nhận.",
    "interval_minutes": "Chu kỳ nguồn 15/30/60 phút; phải khớp khai báo MDMS.",
    "date": "Ngày theo định dạng YYYY-MM-DD.",
    "year_month": "Tháng theo định dạng YYYY-MM.",
    "billing_energy": "Tổng điện năng thương phẩm toàn bộ ngành trong tháng, không chỉ tập mẫu.",
    "billing_basis": "calendar_month hoặc billing_label_month; khai báo theo kỳ thực tế.",
    "organization_scope": "Phạm vi của tổng thương phẩm, khớp phạm vi nghiên cứu.",
    "source_description": "Nguồn và cơ sở xác nhận tổng thương phẩm.",
    "day_type": "weekday / saturday / sunday / holiday; đây là ví dụ, không phải lịch lễ chính thức.",
}


@dataclass(frozen=True)
class InputTemplate:
    title: str
    category: str
    examples: tuple[dict, ...]
    guidance: str

    @property
    def columns(self) -> tuple[str, ...]:
        return tuple(self.examples[0])

    def column_help(self, name: str) -> str:
        if name.startswith("slot_"):
            slots = sum(column.startswith("slot_") for column in self.columns)
            index = int(name.removeprefix("slot_")) - 1
            start, end = index * 1440 // slots, (index + 1) * 1440 // slots
            return (
                f"kWh trong chu kỳ {start // 60:02}:{start % 60:02}–"
                f"{end // 60:02}:{end % 60:02}; thiếu để trống, 0 là số đo hợp lệ."
            )
        return FIELD_HELP[name]


TEMPLATES = {
    "cmis": InputTemplate(
        "CMIS · Khách hàng mẫu",
        "cmis",
        tuple(
            dict(
                customer_id=f"{n:06}",
                industry_code="01",
                industry_name="Cơ khí",
                meter_id=f"M{n:06}",
                organization_code="NPC",
                sampling_weight=1,
            )
            for n in (1, 2)
        ),
        "Tập mẫu đã chốt. Một hàng cho mỗi quan hệ khách hàng–công tơ. "
        "Nếu thay/cộng gộp công tơ, bổ sung valid_from, valid_to, meter_role theo hướng dẫn đầu vào.",
    ),
    "mdms_long": InputTemplate(
        "MDMS · Dạng dài, 30 phút",
        "mdms",
        tuple(
            dict(
                meter_id=f"M{n:06}",
                timestamp="2026-01-01T00:00:00",
                value=n,
                source_quality_flag="actual",
                unit="kWh",
                direction="import",
                interval_minutes=30,
            )
            for n in (1, 2)
        ),
        "Mỗi hàng là một chu kỳ công tơ. Ví dụ: kWh đã áp hệ số, dấu đầu chu kỳ, 30 phút. "
        "Khai báo đúng ngữ nghĩa từng nguồn MDMS trước khi chạy; hai dòng ví dụ chưa đủ một tháng.",
    ),
    "billing": InputTemplate(
        "Thương phẩm · Tổng điện năng ngành-tháng",
        "billing",
        (
            dict(
                industry_code="01",
                year_month="2026-01",
                billing_energy=100000,
                unit="kWh",
                billing_basis="calendar_month",
                organization_scope="NPC",
                source_description="Tổng minh hoạ, thay bằng số đã xác nhận",
            ),
        ),
        "Mỗi ngành-tháng có đúng một tổng toàn ngành. Không lấy điện năng của tập mẫu làm tổng này.",
    ),
    "calendar": InputTemplate(
        "Lịch · Ngày đặc biệt / ghi đè",
        "calendar",
        (
            dict(date="2026-01-01", day_type="holiday"),
            dict(date="2026-01-17", day_type="weekday"),
        ),
        "Tuỳ chọn. Chỉ khai báo ngày cần ghi đè. Ví dụ không phải lịch lễ đã xác nhận; "
        "hãy thay bằng lịch đã chốt của đợt nghiên cứu.",
    ),
}
for _slots in (24, 48, 96):
    TEMPLATES[f"mdms_wide_{_slots}"] = InputTemplate(
        f"MDMS · Dạng rộng, {_slots} chu kỳ/ngày",
        "mdms",
        tuple(
            dict(
                meter_id=f"M{n:06}",
                date="2026-01-01",
                source_quality_flag="actual",
                unit="kWh",
                direction="import",
                interval_minutes=1440 // _slots,
                **{f"slot_{s:02}": n * 48 / _slots for s in range(1, _slots + 1)},
            )
            for n in (1, 2)
        ),
        f"Một hàng công tơ/ngày, đủ {_slots} cột slot_01…slot_{_slots:02}. "
        f"Khai báo dạng rộng, {1440 // _slots} phút, kWh đã áp hệ số, dấu đầu chu kỳ. "
        "Các mẫu dài/rộng là phương án thay thế; không nạp cùng số đo nhiều lần.",
    )


def _write_excel(template: InputTemplate, path: Path) -> None:
    book = openpyxl.Workbook()
    data = book.active
    data.title = "Dữ liệu"
    data.append(template.columns)
    for row in template.examples:
        data.append([row[name] for name in template.columns])
    data.freeze_panes = "C2" if template.category == "mdms" else "A2"
    data.auto_filter.ref = data.dimensions
    guide = book.create_sheet("Hướng dẫn")
    guide.append(["Nội dung", "Hướng dẫn"])
    guide.append(["Lưu ý", NOTICE])
    guide.append([template.title, template.guidance])
    for name in template.columns:
        guide.append([name, template.column_help(name)])
    for sheet in (data, guide):
        sheet.sheet_view.showGridLines = False
        for row in sheet:
            for cell in row:
                cell.font = Font(name="Calibri", size=11, color="172B4D")
                cell.alignment = Alignment(vertical="center", wrap_text=True)
                if isinstance(cell.value, str):
                    cell.number_format = "@"
            sheet.row_dimensions[row[0].row].height = 36 if sheet is data else 48
        for cell in sheet[1]:
            cell.fill = PatternFill("solid", fgColor="173D5D")
            cell.font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
        for index in range(1, sheet.max_column + 1):
            sheet.column_dimensions[get_column_letter(index)].width = 26
    guide.column_dimensions["A"].width = 34
    guide.column_dimensions["B"].width = 100
    guide.freeze_panes = "A2"
    try:
        book.save(path)
    finally:
        book.close()


def write_template(key: str, path: Path) -> Path:
    path = Path(path)
    if key not in TEMPLATES or path.suffix.lower() not in (".csv", ".xlsx"):
        raise ValueError("Chọn mẫu hợp lệ và định dạng CSV hoặc XLSX.")
    if path.exists() or path.is_symlink():
        raise FileExistsError(
            "Tệp đã tồn tại. Hãy chọn tên khác để giữ nguyên dữ liệu."
        )
    path.parent.mkdir(parents=True, exist_ok=True)
    template = TEMPLATES[key]
    with tempfile.TemporaryDirectory(prefix=".template-", dir=path.parent) as staging:
        temporary = Path(staging) / path.name
        if path.suffix.lower() == ".xlsx":
            _write_excel(template, temporary)
        else:
            with temporary.open("w", encoding="utf-8-sig", newline="") as stream:
                writer = csv.DictWriter(stream, fieldnames=template.columns)
                writer.writeheader()
                writer.writerows(template.examples)
        os.link(temporary, path)
    return path
