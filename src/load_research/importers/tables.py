"""Streaming table readers; source values are data, never executable code."""

import csv
import itertools
import zipfile
from collections.abc import Generator, Iterator
from datetime import date, datetime
from pathlib import Path

import openpyxl
import pyarrow.parquet as parquet


def _text(value: object) -> str:
    if value is None:
        return ""
    if isinstance(value, (date, datetime)):
        return value.isoformat()
    return str(value)


def _excel_rows(path: Path, sheet: str | None) -> Iterator[list]:
    with zipfile.ZipFile(path) as archive:
        members = archive.infolist()
        if (
            len(members) > 10000
            or sum(item.file_size for item in members) > 2 * 1024**3
        ):
            raise ValueError(
                "Excel archive exceeds safe expansion limit; use CSV/Parquet"
            )
        if any(
            item.file_size > 1024**2
            and item.file_size > 1000 * max(1, item.compress_size)
            for item in members
        ):
            raise ValueError("Excel archive has an unsafe compression ratio")
    values = openpyxl.load_workbook(
        path, read_only=True, data_only=True, keep_links=False
    )
    formulas = openpyxl.load_workbook(
        path, read_only=True, data_only=False, keep_links=False
    )
    try:
        name = sheet or values.sheetnames[0]
        if name not in values.sheetnames:
            raise ValueError("requested Excel sheet does not exist")
        for cached_row, source_row in zip(
            values[name].iter_rows(), formulas[name].iter_rows(), strict=True
        ):
            for cached, source in zip(cached_row, source_row, strict=True):
                if source.data_type == "f" and cached.value is None:
                    raise ValueError(
                        "Excel formula has no cached value; export values before import"
                    )
            yield [cell.value for cell in cached_row]
    finally:
        values.close()
        formulas.close()


def _raw_rows(path: Path, sheet: str | None, delimiter: str) -> Iterator[list]:
    suffix = path.suffix.lower()
    if suffix == ".csv":
        with path.open(encoding="utf-8-sig", newline="") as stream:
            yield from csv.reader(stream, delimiter=delimiter, strict=True)
    elif suffix == ".xlsx":
        yield from _excel_rows(path, sheet)
    elif suffix == ".parquet":
        source = parquet.ParquetFile(path)
        yield source.schema_arrow.names
        for batch in source.iter_batches(batch_size=20000):
            columns = batch.to_pydict()
            yield from zip(
                *(columns[name] for name in source.schema_arrow.names), strict=True
            )
    else:
        raise ValueError("unsupported input format; expected CSV, XLSX or Parquet")


def read_rows(
    path: Path,
    *,
    mapping: dict[str, str] | None = None,
    sheet: str | None = None,
    delimiter: str = ",",
    header_row: int = 1,
) -> Generator[dict, None, None]:
    path = Path(path)
    if not path.is_file() or len(delimiter) != 1 or header_row < 1:
        raise ValueError("invalid input file or table options")
    rows = iter(_raw_rows(path, sheet, delimiter))
    try:
        for _ in range(header_row - 1):
            next(rows)
        headers = [_text(value).strip() for value in next(rows)]
    except StopIteration as error:
        raise ValueError("table has no header") from error
    if (
        not headers
        or any(not header for header in headers)
        or len(set(headers)) != len(headers)
    ):
        raise ValueError("ambiguous or empty table header")
    aliases = mapping or {name: name for name in headers}
    if any(source not in headers for source in aliases.values()):
        raise ValueError("mapping references a missing source column")
    reserved = {"_source_row", "_source_path", "_source_sheet"}
    if reserved.intersection(aliases):
        raise ValueError("mapping uses reserved provenance fields")
    for number, values in enumerate(rows, start=header_row + 1):
        if not values or all(value is None or value == "" for value in values):
            continue
        if len(values) != len(headers):
            raise ValueError(f"column count differs from header at source row {number}")
        row = dict(zip(headers, (_text(value) for value in values), strict=True))
        mapped: dict = {canonical: row[source] for canonical, source in aliases.items()}
        mapped.update(
            {"_source_row": number, "_source_path": str(path), "_source_sheet": sheet}
        )
        yield mapped


def read_batches(
    path: Path, *, batch_size: int = 20000, **options
) -> Iterator[list[dict]]:
    if type(batch_size) is not int or not 1 <= batch_size <= 1_000_000:
        raise ValueError("batch_size must be between 1 and 1000000")
    rows = iter(read_rows(path, **options))
    while batch := list(itertools.islice(rows, batch_size)):
        yield batch
